"""Load contract text from .txt, .md, .docx, and .pdf files."""

import os

SUPPORTED_EXTENSIONS = (".txt", ".md", ".docx", ".pdf")


def collect_paths(inputs):
    """Expand files and directories into supported document paths.

    Args:
        inputs: File and folder paths from the command line.

    Returns:
        Sorted, de-duplicated paths to .txt, .md, .docx, and .pdf files.

    Raises:
        FileNotFoundError: If an input does not exist.
        ValueError: If an explicit file has an unsupported extension.
    """
    paths = []
    for item in inputs:
        if os.path.isdir(item):
            for root, _dirs, files in os.walk(item):
                for name in files:
                    if name.lower().endswith(SUPPORTED_EXTENSIONS) and not name.startswith("~$"):
                        paths.append(os.path.join(root, name))
        elif os.path.isfile(item):
            if not item.lower().endswith(SUPPORTED_EXTENSIONS):
                raise ValueError(f"Unsupported file type: {item} "
                                 f"(expected {', '.join(SUPPORTED_EXTENSIONS)})")
            paths.append(item)
        else:
            raise FileNotFoundError(f"No such file or directory: {item}")
    return sorted(dict.fromkeys(paths))


def read_docx(path):
    """Return paragraph and table text from a .docx file."""
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("Reading .docx needs python-docx: pip install -r requirements.txt") from exc
    document = docx.Document(path)
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def read_pdf(path):
    """Return extracted text from a .pdf file."""
    try:
        import pypdf
    except ImportError as exc:
        raise RuntimeError("Reading .pdf needs pypdf: pip install -r requirements.txt") from exc
    reader = pypdf.PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def load_text(path):
    """Return the plain text of a supported document.

    Args:
        path: Path to a .txt, .md, .docx, or .pdf file.

    Returns:
        The document text, whitespace-trimmed.

    Raises:
        ValueError: If no text can be extracted.
        RuntimeError: If the optional .docx or .pdf parser is missing.
    """
    lower = path.lower()
    if lower.endswith(".docx"):
        text = read_docx(path)
    elif lower.endswith(".pdf"):
        text = read_pdf(path)
    else:
        with open(path, encoding="utf-8", errors="replace") as handle:
            text = handle.read()
    text = text.strip()
    if not text:
        raise ValueError(f"No extractable text in {path} (scanned PDFs need OCR first)")
    return text
